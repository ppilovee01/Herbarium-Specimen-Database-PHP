import sys
import os
import re
import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_formatted_text(paragraph, text, font_name='TH Sarabun New', size_pt=16):
    # Remove markdown link formatting e.g. [text](url) -> text
    text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', text)
    
    parts = text.split('**')
    for idx, part in enumerate(parts):
        is_bold = (idx % 2 == 1)
        subparts = part.split('`')
        for sidx, subpart in enumerate(subparts):
            is_code = (sidx % 2 == 1)
            if not subpart:
                continue
            run = paragraph.add_run(subpart)
            apply_font_to_run(run, font_name)
            run.font.size = Pt(size_pt)
            if is_bold:
                run.bold = True
            if is_code:
                run.font.name = 'Courier New'
                run.font.size = Pt(11.5)

def apply_font_to_run(run, font_name='TH Sarabun New'):
    run.font.name = font_name
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rPr.append(rFonts)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        '<w:tblBorders %s>'
        '  <w:top w:val="single" w:sz="6" w:space="0" w:color="888888"/>'
        '  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="888888"/>'
        '  <w:left w:val="none"/>'
        '  <w:right w:val="none"/>'
        '  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="DDDDDD"/>'
        '  <w:insideV w:val="none"/>'
        '</w:tblBorders>' % nsdecls('w')
    )
    tblPr.append(borders)

def add_page_number(run):
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    
    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)

def add_cover_page(doc, title, subtitle):
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(80)
    
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(12)
    run_title = p_title.add_run(title)
    run_title.bold = True
    run_title.font.size = Pt(28)
    apply_font_to_run(run_title, 'TH Sarabun New')
    
    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(160)
    run_sub = p_sub.add_run(subtitle)
    run_sub.font.size = Pt(18)
    apply_font_to_run(run_sub, 'TH Sarabun New')
    
    # Presentation info
    p_info = doc.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_info.paragraph_format.line_spacing = 1.3
    
    info_text = (
        "เสนอ\n"
        "อาจารย์ผู้สอนรายวิชาการเขียนโปรแกรมบนเว็บ\n\n\n\n"
        "จัดทำโดย\n"
        "คณะผู้จัดทำโครงงานคอมพิวเตอร์\n"
        ".........................................................................\n"
        ".........................................................................\n\n\n\n"
        "รายงานนี้เป็นส่วนหนึ่งของโครงงานวิชาการเขียนโปรแกรมบนเว็บ (Web Programming)\n"
        "ภาคการศึกษาที่ 3 ปีการศึกษา 2568\n"
        "มหาวิทยาลัยเทคโนโลยีและสารสนเทศ"
    )
    
    run_info = p_info.add_run(info_text)
    run_info.font.size = Pt(16)
    apply_font_to_run(run_info, 'TH Sarabun New')
    
    doc.add_page_break()

def markdown_to_docx(md_path, docx_path, cover_title, cover_sub):
    if not os.path.exists(md_path):
        print(f"Error: {md_path} not found")
        return

    doc = Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25) # Slightly wider left margin for binding
        section.right_margin = Inches(1.0)
        section.different_first_page_header_footer = True
        
        # Setup page number in footer
        footer = section.footer
        footer_p = footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer_run = footer_p.add_run("หน้า ")
        apply_font_to_run(footer_run, 'TH Sarabun New')
        footer_run.font.size = Pt(12)
        add_page_number(footer_p.add_run())

    # Set normal style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'TH Sarabun New'
    font.size = Pt(16) # Standard body text size
    
    rPr = style.element.rPr
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        style.element.append(rPr)
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'TH Sarabun New')
    rFonts.set(qn('w:hAnsi'), 'TH Sarabun New')
    rFonts.set(qn('w:cs'), 'TH Sarabun New')
    rPr.append(rFonts)

    # Add Cover Page
    add_cover_page(doc, cover_title, cover_sub)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.read().split('\n')

    in_code_block = False
    code_text = []
    in_table = False
    table_rows = []
    first_heading = True

    i = 0
    while i < len(lines):
        line = lines[i]
        
        # Code Block
        if line.strip().startswith('```'):
            if in_code_block:
                in_code_block = False
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(6)
                
                # Left accent border
                pBdr = OxmlElement('w:pBdr')
                left_border = OxmlElement('w:left')
                left_border.set(qn('w:val'), 'single')
                left_border.set(qn('w:sz'), '18') # 2.25pt
                left_border.set(qn('w:space'), '6')
                left_border.set(qn('w:color'), '777777')
                pBdr.append(left_border)
                p._p.get_or_add_pPr().append(pBdr)
                
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F5F6F8"/>')
                p._p.get_or_add_pPr().append(shd)

                code_content = '\n'.join(code_text)
                run = p.add_run(code_content)
                run.font.name = 'Courier New'
                run.font.size = Pt(11.0)
                code_text = []
            else:
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_text.append(line)
            i += 1
            continue

        # Table
        if line.strip().startswith('|'):
            in_table = True
            table_rows.append(line)
            i += 1
            continue
        elif in_table:
            in_table = False
            parsed_rows = []
            for tr in table_rows:
                cells = [c.strip() for c in tr.split('|')[1:-1]]
                parsed_rows.append(cells)
            
            filtered_rows = []
            for r in parsed_rows:
                if len(r) > 0 and all(re.match(r'^:?-+:?$', cell) for cell in r if cell):
                    continue
                filtered_rows.append(r)

            if len(filtered_rows) > 0:
                cols_count = len(filtered_rows[0])
                table = doc.add_table(rows=len(filtered_rows), cols=cols_count)
                set_table_borders(table)
                table.autofit = True
                
                for r_idx, row_data in enumerate(filtered_rows):
                    row = table.rows[r_idx]
                    is_header = (r_idx == 0)
                    
                    for c_idx, cell_data in enumerate(row_data):
                        if c_idx >= len(row.cells):
                            break
                        cell = row.cells[c_idx]
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(4)
                        p.paragraph_format.space_after = Pt(4)
                        
                        if is_header:
                            run = p.add_run(cell_data)
                            run.bold = True
                            apply_font_to_run(run)
                            run.font.size = Pt(15)
                            set_cell_background(cell, "365F91") # Formal dark blue
                            run.font.color.rgb = RGBColor(255, 255, 255)
                        else:
                            add_formatted_text(p, cell_data, size_pt=14)
                            
                            if r_idx % 2 == 1:
                                set_cell_background(cell, "F2F5F8")
                            else:
                                set_cell_background(cell, "FFFFFF")
                                
                        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                
                doc.add_paragraph()
            table_rows = []

        # Headings
        if line.startswith('# '):
            # Page break on Level 1 Headings (except the first one after the cover page)
            if not first_heading:
                doc.add_page_break()
            else:
                first_heading = False
                
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(20)
            p.paragraph_format.space_after = Pt(10)
            run = p.add_run(line[2:])
            run.bold = True
            run.font.size = Pt(22)
            run.font.color.rgb = RGBColor(31, 78, 121)
            apply_font_to_run(run)
        elif line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(line[3:])
            run.bold = True
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(56, 124, 186)
            apply_font_to_run(run)
        elif line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(line[4:])
            run.bold = True
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(46, 117, 89)
            apply_font_to_run(run)
        elif line.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[5:])
            run.bold = True
            run.font.size = Pt(16)
            apply_font_to_run(run)
        # Bullet list
        elif line.strip().startswith('- ') or line.strip().startswith('* '):
            indent_level = len(line) - len(line.lstrip())
            text = line.strip()[2:]
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.left_indent = Inches(0.25 + 0.25 * (indent_level // 2))
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.space_before = Pt(0)
            add_formatted_text(p, text, size_pt=16)
        # Numbered list
        elif re.match(r'^\d+\.\s', line.strip()):
            indent_level = len(line) - len(line.lstrip())
            match = re.match(r'^(\d+)\.\s(.*)', line.strip())
            text = match.group(2)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.left_indent = Inches(0.25 + 0.25 * (indent_level // 2))
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.space_before = Pt(0)
            add_formatted_text(p, text, size_pt=16)
        # Horizontal Rule
        elif line.strip() == '---':
            # Skip horizontal rules for clean academic layout
            pass
        # Empty Line
        elif not line.strip():
            pass
        # Normal Paragraph
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_text(p, line, size_pt=16)
                
        i += 1

    doc.save(docx_path)
    print(f"Successfully converted {md_path} to {docx_path}")

if __name__ == '__main__':
    if len(sys.argv) < 5:
        # Defaults
        markdown_to_docx(
            'system_analysis_and_design.md', 
            'system_analysis_and_design.docx',
            'รายงานการวิเคราะห์และออกแบบระบบ',
            'ระบบจัดการและสืบค้นฐานข้อมูลพรรณไม้แห้ง\n(Herbarium Web Application)'
        )
        markdown_to_docx(
            'system_development_results.md', 
            'system_development_results.docx',
            'รายงานการพัฒนาและผลการดำเนินงาน',
            'ระบบจัดการและสืบค้นฐานข้อมูลพรรณไม้แห้ง\n(Herbarium Web Application)'
        )
    else:
        markdown_to_docx(sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4])
